from django.utils.translation import gettext as _
from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db.models import Min
import pandas as pd
from datetime import datetime
from io import BytesIO
from .models import SalesData, PromotionData, ProductData
import numpy as np
import re
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from urllib.parse import urlencode
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def _parse_date(val):
    if pd.isna(val):
        return None
    if isinstance(val, (pd.Timestamp, datetime)):
        return val.date()
    try:
        return datetime.strptime(str(val), '%Y-%m-%d').date()
    except Exception:
        try:
            return pd.to_datetime(val).date()
        except Exception:
            return None


def _normalize_header(value):
    if pd.isna(value):
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    header = str(value)
    if header.endswith('.0') and header[:-2].isdigit():
        header = header[:-2]
    header = header.replace('\u3000', ' ').replace('\xa0', ' ')
    header = re.sub(r'\s+', ' ', header)
    return header.strip()

@login_required
def download_sales_template(request):
    """下载销售数据模板"""
    # 长表格格式：每行包含一个日期和对应销量
    template_data = {
        '销售区域代码': ['CN-SH', 'CN-BJ'],
        '销售区域名称': ['上海', '北京'],
        '产品一级分类': ['饮料', '饮料'],
        '产品二级分类': ['碳酸饮料', '碳酸饮料'],
        '产品三级分类': ['可乐', '可乐'],
        '产品四级分类': ['包装', '大包装'],
        '产品五级分类': ['', ''],
        '物料号': ['M001', 'M002'],
        '物料名称': ['可口可乐', '百事可乐'],
        '日期': ['2023-01-01', '2023-02-01'],
        '销量': [100, 120],
    }

    df = pd.DataFrame(template_data)
    
    # 保存为Excel
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)

    response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="sales_data_template.xlsx"'
    return response

@login_required
def download_promotion_template(request):
    """下载促销数据模板"""
    # 创建模板DataFrame（列顺序按照要求）
    template_data = {
        '销售区域代码': ['CN-SH', 'CN-BJ'],
        '销售区域名称': ['上海', '北京'],
        '产品一级分类': ['饮料', '饮料'],
        '产品二级分类': ['碳酸饮料', '碳酸饮料'],
        '产品三级分类': ['可乐', '可乐'],
        '产品四级分类': ['', ''],
        '产品五级分类': ['', ''],
        '物料号': ['M001', 'M002'],
        '物料名称': ['可口可乐', '百事可乐'],
        '日期': ['2023-01-01', '2023-02-01'],
        '促销力度': [0.2, 0.15],
        '促销开始日期': ['2023-01-05', '2023-02-01'],
        '促销结束日期': ['2023-01-20', '2023-02-10']
    }
    df = pd.DataFrame(template_data)
    
    # 保存为Excel
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)

    response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="promotion_data_template.xlsx"'
    return response

@login_required
def upload_data(request):
    """上传销售或促销 Excel 并保存到数据库，包含基础的数据质量处理。"""
    if request.method == 'POST':
        uploaded_count = 0
        errors = []

        if not request.FILES:
            messages.error(request, _('未检测到上传文件，请选择要上传的销售或促销数据文件。'))
            return redirect('data:upload')

        # iterate over possible file inputs
        for key in ('sales_file', 'promotion_file'):
            if key not in request.FILES:
                continue
            f = request.FILES[key]
            try:
                df = pd.read_excel(f, engine='openpyxl')
            except Exception as e:
                messages.error(request, _(f'无法读取文件 {f.name}: {e}'))
                continue

            # normalize column names to simplify downstream
            df.columns = [_normalize_header(c) for c in df.columns]
            month_columns = [c for c in df.columns if re.fullmatch(r'20\d{4}', str(c))]
            wide_mode = bool(month_columns) and not ('日期' in df.columns and '销量' in df.columns)

            # basic required fields
            if key == 'sales_file':
                if '物料号' not in df.columns or '物料名称' not in df.columns or (not wide_mode and ('日期' not in df.columns or '销量' not in df.columns)):
                    messages.error(request, _(f'{f.name} 缺少必要字段: 物料号/物料名称/日期/销量 或 年月列'))
                    continue
            else:
                if '物料号' not in df.columns or '物料名称' not in df.columns or '日期' not in df.columns or '促销力度' not in df.columns:
                    messages.error(request, _(f'{f.name} 缺少必要字段: 物料号/物料名称/日期/促销力度'))
                    continue

            # fill basic missing values
            subset = ['物料号']
            if '销售区域代码' in df.columns:
                subset.append('销售区域代码')
            if '日期' in df.columns:
                subset.append('日期')
            df = df.drop_duplicates(subset=subset)
            if '销量' in df.columns and not wide_mode:
                df['销量'] = df['销量'].replace([np.nan], 0)
                df.loc[df['销量'] < 0, '销量'] = 0

            for row_index, (_idx, row) in enumerate(df.iterrows()):
                row_number = row_index + 2
                try:
                    material_number_raw = row.get('物料号')
                    material_name_raw = row.get('物料名称')
                    if pd.isna(material_number_raw) or pd.isna(material_name_raw):
                        errors.append(f'第 {row_number} 行缺少物料号或物料名称')
                        continue

                    material_number = str(material_number_raw).strip()
                    material_name = str(material_name_raw).strip()
                    date_val = _parse_date(row.get('日期')) if '日期' in df.columns else None
                    if not wide_mode and not date_val:
                        errors.append(f'第 {row_number} 行日期格式无效: {row.get("日期")}')
                        continue

                    def _clean_text(column_name):
                        raw = row.get(column_name)
                        if pd.isna(raw):
                            return ''
                        return str(raw).strip()

                    prod_kwargs = {
                        'material_number': material_number,
                        'material_name': material_name,
                        'product_category_level1': _clean_text('产品一级分类'),
                        'product_category_level2': _clean_text('产品二级分类'),
                        'product_category_level3': _clean_text('产品三级分类'),
                        'product_category_level4': _clean_text('产品四级分类') if '产品四级分类' in df.columns else '',
                        'product_category_level5': _clean_text('产品五级分类') if '产品五级分类' in df.columns else '',
                    }
                    product, created = ProductData.objects.get_or_create(material_number=material_number, defaults=prod_kwargs)
                    if not created:
                        updated = False
                        for k, v in prod_kwargs.items():
                            if getattr(product, k) != v:
                                setattr(product, k, v)
                                updated = True
                        if updated:
                            product.save()

                    def _clean_attr(column_name):
                        if column_name not in df.columns:
                            return ''
                        raw = row.get(column_name)
                        return str(raw).strip() if pd.notna(raw) else ''

                    region_code = _clean_attr('销售区域代码')
                    region_name = _clean_attr('销售区域名称')

                    if key == 'sales_file':
                        if wide_mode:
                            row_uploaded = 0
                            for month_col in month_columns:
                                qty_val = row.get(month_col)
                                if pd.isna(qty_val):
                                    continue
                                try:
                                    qty = float(qty_val)
                                except (TypeError, ValueError):
                                    qty = 0.0
                                try:
                                    month_date = datetime.strptime(str(month_col), '%Y%m').date()
                                except Exception:
                                    errors.append(f'第 {row_number} 行无效月份列: {month_col}')
                                    continue
                                SalesData.objects.update_or_create(
                                    product_data=product,
                                    sales_region_code=region_code,
                                    date=month_date,
                                    defaults={
                                        'sales_region_name': region_name,
                                        'sales_quantity': qty,
                                    }
                                )
                                uploaded_count += 1
                                row_uploaded += 1
                            if row_uploaded == 0:
                                errors.append(f'第 {row_number} 行未找到有效月度销量数据')
                        else:
                            qty_val = row.get('销量') if '销量' in df.columns else None
                            try:
                                qty = float(qty_val) if qty_val is not None and pd.notna(qty_val) else 0.0
                            except (TypeError, ValueError):
                                qty = 0.0
                            SalesData.objects.update_or_create(
                                product_data=product,
                                sales_region_code=region_code,
                                date=date_val,
                                defaults={
                                    'sales_region_name': region_name,
                                    'sales_quantity': qty,
                                }
                            )
                            uploaded_count += 1
                    else:
                        promo_val = row.get('促销力度') if '促销力度' in df.columns else None
                        try:
                            promo_strength = float(promo_val) if promo_val is not None and pd.notna(promo_val) else None
                        except (TypeError, ValueError):
                            promo_strength = None
                        start_date = _parse_date(row.get('促销开始日期')) if '促销开始日期' in df.columns else None
                        end_date = _parse_date(row.get('促销结束日期')) if '促销结束日期' in df.columns else None
                        PromotionData.objects.update_or_create(
                            product_data=product,
                            sales_region_code=region_code,
                            date=date_val,
                            defaults={
                                'sales_region_name': region_name,
                                'promotion_strength': promo_strength,
                                'promotion_start_date': start_date,
                                'promotion_end_date': end_date,
                            }
                        )
                        uploaded_count += 1
                except Exception as e:
                    errors.append(f'第 {row_number} 行错误: {e}')

        if uploaded_count:
            messages.success(request, _(f'上传成功: {uploaded_count} 条'))
        if errors:
            messages.error(request, _(f'部分记录处理失败: {len(errors)} 条，示例: {errors[:3]}'))
        elif uploaded_count == 0:
            messages.error(request, _('上传完成，但未检测到有效记录，请检查 Excel 是否符合模板格式。'))

        return redirect('data:view_data')

    return render(request, 'data_management/data_upload.html')

@login_required
def view_data(request):
    """查看数据视图"""
    # 获取筛选参数
    region = request.GET.get('region', '')
    sku = request.GET.get('sku', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')

    if not start_date:
        sales_min = SalesData.objects.aggregate(min_date=Min('date'))['min_date']
        promo_min = PromotionData.objects.aggregate(min_date=Min('date'))['min_date']
        earliest = None
        if sales_min and promo_min:
            earliest = sales_min if sales_min <= promo_min else promo_min
        else:
            earliest = sales_min or promo_min
        if earliest:
            start_date = earliest.isoformat()
    
    # 查询销售数据
    sales_qs = SalesData.objects.select_related('product_data').all()
    if region:
        sales_qs = sales_qs.filter(sales_region_code=region)
    if sku:
        sales_qs = sales_qs.filter(product_data__material_number=sku)
    if start_date:
        sales_qs = sales_qs.filter(date__gte=start_date)
    if end_date:
        sales_qs = sales_qs.filter(date__lte=end_date)

    # 排序
    sales_order = request.GET.get('sales_order', '-date')
    if sales_order:
        try:
            sales_qs = sales_qs.order_by(sales_order)
        except Exception:
            sales_qs = sales_qs.order_by('-date')

    # 分页（每页30条）
    sales_paginator = Paginator(sales_qs, 30)
    sales_page = request.GET.get('sales_page', 1)
    try:
        sales_data = sales_paginator.page(sales_page)
    except PageNotAnInteger:
        sales_data = sales_paginator.page(1)
    except EmptyPage:
        sales_data = sales_paginator.page(sales_paginator.num_pages)
    # 计算限宽页码（最多显示5个）
    def _page_numbers(page_obj, window=5):
        total = page_obj.paginator.num_pages
        current = page_obj.number
        half = window // 2
        start = max(1, current - half)
        end = min(total, start + window - 1)
        start = max(1, end - window + 1)
        return list(range(start, end + 1))
    sales_page_numbers = _page_numbers(sales_data, window=5)

    # 查询促销数据
    promo_qs = PromotionData.objects.select_related('product_data').all()
    if region:
        promo_qs = promo_qs.filter(sales_region_code=region)
    if sku:
        promo_qs = promo_qs.filter(product_data__material_number=sku)
    if start_date:
        promo_qs = promo_qs.filter(date__gte=start_date)
    if end_date:
        promo_qs = promo_qs.filter(date__lte=end_date)

    promo_order = request.GET.get('promo_order', '-date')
    if promo_order:
        try:
            promo_qs = promo_qs.order_by(promo_order)
        except Exception:
            promo_qs = promo_qs.order_by('-date')

    promo_paginator = Paginator(promo_qs, 30)
    promo_page = request.GET.get('promo_page', 1)
    try:
        promotion_data = promo_paginator.page(promo_page)
    except PageNotAnInteger:
        promotion_data = promo_paginator.page(1)
    except EmptyPage:
        promotion_data = promo_paginator.page(promo_paginator.num_pages)
    promo_page_numbers = _page_numbers(promotion_data, window=5)

    # 获取唯一区域和SKU用于筛选
    sales_regions = list(SalesData.objects.values_list('sales_region_code', flat=True).distinct())
    promo_regions = list(PromotionData.objects.values_list('sales_region_code', flat=True).distinct())
    regions = sorted(set([r for r in sales_regions if r] + [r for r in promo_regions if r]))
    skus = ProductData.objects.values_list('material_number', flat=True).distinct()

    # 构造不含分页和排序的基础查询字符串，用于在模板生成链接时复用
    exclude_keys = {'sales_page', 'promo_page', 'sales_order', 'promo_order'}
    base_params = {k: v for k, v in request.GET.items() if k not in exclude_keys and v}
    base_qs = urlencode(base_params)

    context = {
        'sales_data': sales_data,
        'promotion_data': promotion_data,
        'regions': regions,
        'skus': skus,
        'selected_region': region,
        'selected_sku': sku,
        'start_date': start_date,
        'end_date': end_date,
        'base_qs': base_qs,
        'current_sales_order': sales_order,
        'current_promo_order': promo_order,
        'sales_page_numbers': sales_page_numbers,
        'promo_page_numbers': promo_page_numbers,
    }

    return render(request, 'data_management/view_data.html', context)

def _get_filtered_data(request):
    region = request.GET.get('region', '')
    sku = request.GET.get('sku', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')

    sales_qs = SalesData.objects.select_related('product_data').all()
    promo_qs = PromotionData.objects.select_related('product_data').all()
    if region:
        sales_qs = sales_qs.filter(sales_region_code=region)
        promo_qs = promo_qs.filter(sales_region_code=region)
    if sku:
        sales_qs = sales_qs.filter(product_data__material_number=sku)
        promo_qs = promo_qs.filter(product_data__material_number=sku)
    if start_date:
        sales_qs = sales_qs.filter(date__gte=start_date)
        promo_qs = promo_qs.filter(date__gte=start_date)
    if end_date:
        sales_qs = sales_qs.filter(date__lte=end_date)
        promo_qs = promo_qs.filter(date__lte=end_date)

    return {
        'sales_qs': sales_qs,
        'promo_qs': promo_qs,
        'selected_region': region,
        'selected_sku': sku,
        'start_date': start_date,
        'end_date': end_date,
    }

@login_required
def export_data(request):
    filtered = _get_filtered_data(request)
    # 支持根据请求参数导出单表或全部
    sheet = request.GET.get('sheet', 'all')
    sales_order = request.GET.get('sales_order', '-date')
    promo_order = request.GET.get('promo_order', '-date')

    try:
        sales_qs = filtered['sales_qs'].order_by(sales_order)
    except Exception:
        sales_qs = filtered['sales_qs'].order_by('-date')
    try:
        promo_qs = filtered['promo_qs'].order_by(promo_order)
    except Exception:
        promo_qs = filtered['promo_qs'].order_by('-date')

    sales_rows = []
    for item in sales_qs:
        sales_rows.append({
            '日期': item.date,
            '销售区域代码': item.sales_region_code,
            '销售区域名称': item.sales_region_name or '',
            '物料号': item.product_data.material_number,
            '物料名称': item.product_data.material_name,
            '产品一级分类': item.product_data.product_category_level1,
            '产品二级分类': item.product_data.product_category_level2,
            '产品三级分类': item.product_data.product_category_level3,
            '产品四级分类': item.product_data.product_category_level4,
            '产品五级分类': item.product_data.product_category_level5,
            '销量': item.sales_quantity,
        })

    promo_rows = []
    for item in promo_qs:
        promo_rows.append({
            '日期': item.date,
            '销售区域代码': item.sales_region_code,
            '销售区域名称': item.sales_region_name or '',
            '物料号': item.product_data.material_number,
            '物料名称': item.product_data.material_name,
            '产品一级分类': item.product_data.product_category_level1,
            '产品二级分类': item.product_data.product_category_level2,
            '产品三级分类': item.product_data.product_category_level3,
            '产品四级分类': item.product_data.product_category_level4,
            '产品五级分类': item.product_data.product_category_level5,
            '促销力度': item.promotion_strength,
            '促销开始日期': item.promotion_start_date,
            '促销结束日期': item.promotion_end_date,
        })

    sales_df = pd.DataFrame(sales_rows)
    promo_df = pd.DataFrame(promo_rows)

    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        workbook = None
        sheets_to_format = []
        if sheet in ('all', 'sales'):
            sales_df.to_excel(writer, index=False, sheet_name='销售数据')
            sheets_to_format.append('销售数据')
        if sheet in ('all', 'promo'):
            promo_df.to_excel(writer, index=False, sheet_name='促销数据')
            sheets_to_format.append('促销数据')

        workbook = writer.book
        for sheet_name in sheets_to_format:
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows(min_row=1, max_row=1):
                for cell in row:
                    cell.font = Font(size=12, bold=True)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                for cell in row:
                    cell.font = Font(size=11)
                    cell.alignment = Alignment(horizontal='left', vertical='center')
            for column_cells in worksheet.columns:
                max_length = max(
                    (len(str(cell.value)) if cell.value is not None else 0)
                    for cell in column_cells
                )
                adjusted_width = max_length + 4
                column_letter = get_column_letter(column_cells[0].column)
                worksheet.column_dimensions[column_letter].width = adjusted_width

    output.seek(0)
    filename = 'data_export.xlsx'
    if sheet == 'sales':
        filename = 'sales_export.xlsx'
    elif sheet == 'promo':
        filename = 'promo_export.xlsx'
    response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

@login_required
def adjust_data(request):
    """数据调整视图"""
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            action = data.get('action')
            
            if action == 'update_sales':
                record_id = data.get('id')
                quantity = data.get('quantity')
                SalesData.objects.filter(id=record_id).update(sales_quantity=quantity)
                return JsonResponse({'success': True, 'message': _('销售数据更新成功')})
            
            elif action == 'update_promotion':
                record_id = data.get('id')
                promotion_strength = data.get('promotion_strength')
                promotion_start_date = _parse_date(data.get('promotion_start_date'))
                promotion_end_date = _parse_date(data.get('promotion_end_date'))
                update_values = {'promotion_strength': promotion_strength}
                if promotion_start_date is not None:
                    update_values['promotion_start_date'] = promotion_start_date
                if promotion_end_date is not None:
                    update_values['promotion_end_date'] = promotion_end_date
                PromotionData.objects.filter(id=record_id).update(**update_values)
                return JsonResponse({'success': True, 'message': _('促销数据更新成功')})
            
            elif action == 'delete_sales':
                record_id = data.get('id')
                SalesData.objects.filter(id=record_id).delete()
                return JsonResponse({'success': True, 'message': _('销售数据删除成功')})
            
            elif action == 'delete_promotion':
                record_id = data.get('id')
                PromotionData.objects.filter(id=record_id).delete()
                return JsonResponse({'success': True, 'message': _('促销数据删除成功')})
                
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=400)
    
    # GET请求 - 显示数据
    sales_page = request.GET.get('sales_page', 1)
    promo_page = request.GET.get('promo_page', 1)

    sales_paginator = Paginator(SalesData.objects.all().order_by('-date'), 30)
    promo_paginator = Paginator(PromotionData.objects.all().order_by('-date'), 30)

    try:
        sales_data = sales_paginator.page(sales_page)
    except PageNotAnInteger:
        sales_data = sales_paginator.page(1)
    except EmptyPage:
        sales_data = sales_paginator.page(sales_paginator.num_pages)

    try:
        promotion_data = promo_paginator.page(promo_page)
    except PageNotAnInteger:
        promotion_data = promo_paginator.page(1)
    except EmptyPage:
        promotion_data = promo_paginator.page(promo_paginator.num_pages)

    def _page_numbers(page_obj, window=5):
        total = page_obj.paginator.num_pages
        current = page_obj.number
        half = window // 2
        start = max(1, current - half)
        end = min(total, start + window - 1)
        start = max(1, end - window + 1)
        return list(range(start, end + 1))

    sales_page_numbers = _page_numbers(sales_data, window=5)
    promo_page_numbers = _page_numbers(promotion_data, window=5)

    return render(request, 'data_management/adjust_data.html', {
        'sales_data': sales_data,
        'promotion_data': promotion_data,
        'sales_page_numbers': sales_page_numbers,
        'promo_page_numbers': promo_page_numbers,
    })
